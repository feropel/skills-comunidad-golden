# -*- coding: utf-8 -*-
"""
AUTOPRUEBA DEL MOTOR — skill golden-dropi-analisis

Siembra exports de Dropi con DEFECTOS CONOCIDOS, corre el motor contra ellos y compara cada
cifra con la verdad calculada a mano. Cada caso nacio de un fallo REAL que inflaba o escondia
cifras de negocio, asi que esto es una red contra regresiones, no una formalidad.

    python3 autoprueba_motor.py            # usa el motor de al lado
    python3 autoprueba_motor.py --mutar    # ademas rompe el motor a proposito y comprueba
                                           # que la autoprueba lo caza (prueba del instrumento)

Sale 0 si todo cuadra, 1 si algo falla. No toca nada fuera de su carpeta temporal.
Requiere openpyxl. Si falta reportlab, el caso del PDF se declara OMITIDO, no se da por bueno.
"""
import os, sys, shutil, subprocess, tempfile, json, io

try:
    import openpyxl
    from openpyxl import Workbook
except ImportError:
    sys.exit("Falta 'openpyxl'. Instalalo con: pip install openpyxl")

AQUI  = os.path.dirname(os.path.abspath(__file__))
MOTOR = os.path.join(AQUI, "motor_analisis_dropi.py")   # se reapunta a una COPIA en modo --mutar
PY    = sys.executable

CP = ["ID","FECHA","ESTATUS","TRANSPORTADORA","DEPARTAMENTO DESTINO","CIUDAD DESTINO",
      "NOMBRE CLIENTE","TELÉFONO","DIRECCION","GANANCIA","VALOR FACTURADO",
      "PRECIO FLETE","COSTO DEVOLUCION FLETE","NOVEDAD"]
CD = ["ID","FECHA","ESTATUS","TRANSPORTADORA","DEPARTAMENTO DESTINO","CIUDAD DESTINO",
      "NOMBRE CLIENTE","TELÉFONO","DIRECCION","PRODUCTO","VARIACION","SKU","CANTIDAD","GANANCIA"]

def esc(path, cols, filas):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.append(cols)
    for f in filas: ws.append(f)
    wb.save(path)

def sembrar(raiz):
    """Los 12 defectos sembrados. Cada uno rompio una cifra de verdad alguna vez."""
    A = os.path.join(raiz, "banco", "Cuenta-A")
    ped = [
      [1001,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA GOMEZ","3001110001","CL 1",20000,80000,9000,0,""],
      [1002,"02-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","BETO RIOS","3001110002","CL 2",20000,80000,9000,0,""],
      [1003,"03-05-2026","DEVOLUCION","ENVIA","ANTIOQUIA","MEDELLIN","CARO DIAZ","3001110003","CL 3",0,80000,9000,9000,"NO CONTESTA"],
      # 1) SINIESTRO = paquete perdido -> devolucion (con igualdad exacta caia en transito)
      [1004,"04-05-2026","SINIESTRO","ENVIA","ANTIOQUIA","MEDELLIN","DORA LUZ","3001110004","CL 4",0,80000,9000,9000,"SINIESTRO"],
      # 2) GUIA_ANULADA -> cancelado
      [1005,"05-05-2026","GUIA_ANULADA","ENVIA","ANTIOQUIA","MEDELLIN","ELI PAZ","3001110005","CL 5",0,80000,9000,0,""],
      # 3) REEXPEDICION CON TILDE -> devolucion (el patron sin tilde no casaba)
      [1006,"06-05-2026","REEXPEDICIÓN","ENVIA","ANTIOQUIA","MEDELLIN","FABI SOL","3001110006","CL 6",0,80000,9000,9000,""],
      # 4) FANTASMA: mismo telefono y monto, un dia aparte, la vieja sin cerrar
      [1007,"07-05-2026","PENDIENTE CONFIRMACION","ENVIA","ANTIOQUIA","MEDELLIN","GINA ORT","3001110007","CL 7",0,150000,9000,0,""],
      [1008,"08-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","GINA ORT","3001110007","CL 7",35000,150000,9000,0,""],
      # 5) transportadora con UNA orden: con umbral desaparecia de la tabla
      [1009,"09-05-2026","ENTREGADO","VELOCES","BOYACA","TUNJA","HUGO MEJIA","3001110009","CL 9",20000,80000,9000,0,""],
      # 6) ciudad con 3 activas: con umbral no daba recomendacion
      [1010,"10-05-2026","ENTREGADO","TCC","NARINO","PASTO","IVAN RUIZ","3001110010","CL 10",20000,80000,9000,0,""],
      [1011,"11-05-2026","DEVOLUCION","TCC","NARINO","PASTO","JOSE LARA","3001110011","CL 11",0,80000,9000,9000,"RECHAZADO"],
      [1012,"12-05-2026","ENTREGADO","TCC","NARINO","PASTO","KARL VEGA","3001110012","CL 12",20000,80000,9000,0,""],
      # 7) cliente REAL con TESTA en el nombre: el filtro por subcadena lo borraba
      [1013,"13-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","MARIA TESTA","3001110013","CL 13",20000,80000,9000,0,""],
      # 8) dinero como TEXTO con decimal de punto: se multiplicaba por 100
      [1014,"14-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","NORA BELL","3001110014","CL 14","20000.00",80000,9000,0,""],
    ]
    esc(os.path.join(A, "2026-05 - por pedido.xlsx"), CP, ped)
    prd = [[r[0],r[1],r[2],r[3],r[4],r[5],r[6],r[7],r[8],"CREMA X","","SKU-1",1,r[9]] for r in ped]
    # 9) dos lineas del MISMO producto con SKU distinto: el dedup se comia una
    prd.append([1001,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA GOMEZ","3001110001","CL 1","CREMA X","ROJO","SKU-2",1,0])
    esc(os.path.join(A, "2026-05 - por producto.xlsx"), CD, prd)
    # 10) y 11) archivos que desaparecian SIN AVISO
    with io.open(os.path.join(A, "2026-05 extra.csv"), "w", encoding="utf-8") as f:
        f.write("ID,FECHA,ESTATUS\n9001,15-05-2026,ENTREGADO\n")
    esc(os.path.join(A, "2026-05 renombrado.xlsx"), ["ID","FECHA","ESTADO","TRANSPORTADORA"],
        [[9002,"16-05-2026","ENTREGADO","ENVIA"]])
    # 12) segunda cuenta: el GLOBAL las suma y eso hay que decirlo
    B = os.path.join(raiz, "banco", "Cuenta-B")
    pb = [[2001,"01-05-2026","ENTREGADO","INTERRAPIDISIMO","VALLE","CALI","LUZ MENA","3002220001","CR 1",50000,200000,9000,0,""],
          [2002,"02-05-2026","DEVOLUCION","INTERRAPIDISIMO","VALLE","CALI","MARIO GIL","3002220002","CR 2",0,200000,9000,9000,"NO CONTESTA"]]
    esc(os.path.join(B, "2026-05 - por pedido.xlsx"), CP, pb)
    esc(os.path.join(B, "2026-05 - por producto.xlsx"), CD,
        [[r[0],r[1],r[2],r[3],r[4],r[5],r[6],r[7],r[8],"SERUM Y","","SKU-9",1,r[9]] for r in pb])

    # --- C1: control POSITIVO. El filtro de prueba TIENE que seguir cazando pruebas de verdad,
    #         y el mismo export bajado dos veces TIENE que salir como duplicado.
    c1 = os.path.join(raiz, "C1", "Cta")
    f1 = [[1,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","CLIENTE REAL UNO","3001110001","CL 1",20000,80000,9000,0,""],
          [2,"02-05-2026","DEVOLUCION","ENVIA","ANTIOQUIA","MEDELLIN","CLIENTE REAL DOS","3001110002","CL 2",0,80000,9000,9000,""],
          [3,"03-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","PRUEBA JUAN","3009999999","CL 3",99999,99999,9000,0,""],
          [4,"04-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","TEST","3009999998","CL 4",88888,88888,9000,0,""]]
    esc(os.path.join(c1, "2026-05 - por pedido.xlsx"), CP, f1)
    esc(os.path.join(c1, "2026-05 - por pedido (1).xlsx"), CP, f1)

    # --- C2: sin informe POR PEDIDO no puede salir un veredicto de rentabilidad
    c2 = os.path.join(raiz, "C2")
    esc(os.path.join(c2, "Cta", "2026-05 - por producto.xlsx"), CD,
        [[10,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1","CREMA","","S1",1,20000]])
    with io.open(os.path.join(c2, "_config_dropi.json"), "w", encoding="utf-8") as f:
        json.dump({"gasto_publicidad": 500000, "negocio": "MI EMPRESA"}, f)

    # --- C3: con config completa, el titulo lleva la marca del CLIENTE y el P&L cuadra
    c3 = os.path.join(raiz, "C3")
    esc(os.path.join(c3, "Cta", "2026-05 - por pedido.xlsx"), CP, [
        [20,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",100000,300000,9000,0,""],
        [21,"02-05-2026","DEVOLUCION","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",0,300000,9000,10000,""]])
    with io.open(os.path.join(c3, "_config_dropi.json"), "w", encoding="utf-8") as f:
        json.dump({"gasto_publicidad": 30000, "negocio": "DISTRIBUIDORA X", "currency": "$"}, f)

    # --- C6: ENTREGADO A TRANSPORTADORA. Es entrega al COURIER: la plata NO ha entrado.
    #         El banco no tenia ni una fila con este estado, y por eso una rotura real pasaba
    #         en verde. Es el estado con la trampa mas cara de todo el universo de Dropi.
    c6 = os.path.join(raiz, "C6", "Cta")
    esc(os.path.join(c6, "2026-05 - por pedido.xlsx"), CP, [
        [60,"01-05-2026","ENTREGADO A TRANSPORTADORA","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",20000,80000,9000,0,""],
        [61,"02-05-2026","ENTREGADO A TRANSPORTADORA","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",20000,80000,9000,0,""],
        [62,"03-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","CARO","3001110003","CL 3",20000,80000,9000,0,""]])

    # --- C7: basura que no puede dar cifras falsas. ID vacio (antes fundia filas y borraba
    #         ventas), dinero en tres formatos que antes salian mal, CANTIDAD 0 y negativa,
    #         nombre de producto en blanco y una fecha en el futuro.
    c7 = os.path.join(raiz, "C7", "Cta")
    esc(os.path.join(c7, "2026-05 - por pedido.xlsx"), CP, [
        ["", "01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",20000,80000,9000,0,""],
        ["", "02-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",20000,80000,9000,0,""],
        [73,"03-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","CARO","3001110003","CL 3","8E+04",80000,9000,0,""],
        [74,"04-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","DORA","3001110004","CL 4","(20.000)",80000,9000,0,""],
        [75,"01-05-2099","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ELI","3001110005","CL 5",10000,80000,9000,0,""]])
    esc(os.path.join(c7, "2026-05 - por producto.xlsx"), CD, [
        [73,"03-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","CARO","3001110003","CL 3","   ","","S1",0,0],
        [74,"04-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","DORA","3001110004","CL 4","CREMA X","","S2",-5,0]])

    # --- C8: solo el export POR PEDIDO. La base de clientes no se puede medir, y ese 0 hay
    #         que decirlo: el PDF remite al usuario a un archivo que sale vacio.
    c8 = os.path.join(raiz, "C8", "Cta")
    esc(os.path.join(c8, "2026-05 - por pedido.xlsx"), CP, [
        [80,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",20000,80000,9000,0,""]])

    # --- C5: TODO en ruta, ninguna orden cerrada. El % de entrega es indefinido, no 0%.
    c5 = os.path.join(raiz, "C5", "Cta")
    esc(os.path.join(c5, "2026-05 - por pedido.xlsx"), CP, [
        [40,"01-05-2026","EN REPARTO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",0,50000,9000,0,""],
        [41,"02-05-2026","PENDIENTE","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",0,50000,9000,0,""]])

    with io.open(os.path.join(raiz, "C5", "_config_dropi.json"), "w", encoding="utf-8") as f:
        json.dump({"gasto_publicidad": 120000, "negocio": "DISTRIBUIDORA X"}, f)

    # --- C4: control NEGATIVO. Compras REALES del mismo cliente: no son fantasmas y no se
    #         pueden declarar como tales. Declarar fantasmas a ciegas ya costo caro una vez.
    c4 = os.path.join(raiz, "C4", "Cta")
    esc(os.path.join(c4, "2026-05 - por pedido.xlsx"), CP, [
        [30,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",20000,80000,9000,0,""],
        [31,"10-06-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","ANA","3001110001","CL 1",20000,95000,9000,0,""],
        [32,"01-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",20000,80000,9000,0,""],
        [33,"03-05-2026","ENTREGADO","ENVIA","ANTIOQUIA","MEDELLIN","BETO","3001110002","CL 2",20000,80000,9000,0,""]])

# --------------------------------------------------------------------------- comprobaciones
class Chequeo:
    def __init__(self): self.fallos = []; self.ok = 0; self.omitidos = []
    def val(self, nombre, real, esperado):
        if str(real) == str(esperado): self.ok += 1; print(f"  OK     {nombre}: {real}")
        else:
            self.fallos.append(f"{nombre}: {real!r} != {esperado!r}")
            print(f"  FALLA  {nombre}: {real!r}  (esperado {esperado!r})")
    def dice(self, nombre, texto, aguja, debe=True):
        if (aguja in texto) == debe: self.ok += 1; print(f"  OK     {nombre}")
        else:
            self.fallos.append(f"{nombre}: {'falta' if debe else 'sobra'} {aguja!r}")
            print(f"  FALLA  {nombre}: {'no aparece' if debe else 'aparece'} {aguja!r}")
    def omitir(self, nombre, porque):
        self.omitidos.append(f"{nombre} ({porque})"); print(f"  OMITIDO {nombre}: {porque}")

def _texto_pdf(ruta):
    """Texto del PDF, para comprobar lo que DICE y no solo que existe. Si no hay extractor
    disponible devuelve None y el caso se declara OMITIDO, nunca por bueno."""
    try:
        from pdfminer.high_level import extract_text
        return extract_text(ruta)
    except Exception:
        pass
    try:
        import subprocess as _sp
        r = _sp.run(["pdftotext", ruta, "-"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip(): return r.stdout
    except Exception:
        pass
    return None

def correr(carpeta):
    r = subprocess.run([PY, MOTOR, carpeta], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr

def celda(xlsx, hoja, etiqueta):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[hoja]
    for row in ws.iter_rows(values_only=True):
        v = [x for x in row if x is not None]
        if v and str(v[0]).strip() == etiqueta: return v[1] if len(v) > 1 else None
    return None

def filas(xlsx, hoja, desde=4):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[hoja]
    return [r[0] for r in ws.iter_rows(min_row=desde, values_only=True) if r and r[0]]

def probar(raiz, c):
    hay_pdf = True
    try: import reportlab  # noqa
    except ImportError: hay_pdf = False

    print("\n=== BANCO: 14 ordenes + 2 de otra cuenta, 12 defectos sembrados ===")
    rc, out = correr(os.path.join(raiz, "banco"))
    c.val("codigo de salida", rc, 0)
    L = os.path.join(raiz, "banco", "Analisis", "MAESTRO_LOGISTICA.xlsx")
    c.val("% entrega GLOBAL", celda(L, "RESUMEN", "% ENTREGA"), "64.3%")
    c.val("entregadas", celda(L, "RESUMEN", "Entregadas"), 9)
    c.val("devoluciones (SINIESTRO y REEXPEDICION dentro)", celda(L, "RESUMEN", "Devoluciones"), 5)
    c.val("canceladas (GUIA_ANULADA dentro)", celda(L, "RESUMEN", "Canceladas/Rechazadas"), 1)
    c.val("ganancia realizada (sin inflar x100)",
          celda(L, "RESUMEN EJECUTIVO", "Ganancia realizada (de lo entregado)"), "$225.000")
    c.dice("declara los 2 archivos saltados", out, "Archivos SALTADOS: 2")
    c.dice("avisa del par fantasma", out, "FANTASMA: 1 par")
    # 🔴 No basta con que el mensaje aparezca: el mensaje se imprime ANTES de cualquier
    # descuento. Lo que hay que medir es que las DOS ordenes sigan contando.
    c.val("NO descuenta la candidata: siguen las 16 ordenes",
          celda(L, "RESUMEN", "Órdenes totales"), 16)
    _wbF = openpyxl.load_workbook(L, data_only=True)["POSIBLES FANTASMA"]
    c.val("y la lista de candidatas trae el par", sum(1 for r in _wbF.iter_rows(min_row=5,
          values_only=True) if r and r[0]), 1)
    c.dice("avisa de que el GLOBAL suma cuentas", out, "Hay 2 cuentas")
    # El aviso de consola se pierde al cerrar el chat: lo que queda es el archivo.
    _cob = openpyxl.load_workbook(L, data_only=True)["COBERTURA"]
    _txt = " ".join(str(v) for r in _cob.iter_rows(values_only=True) for v in r if v)
    c.dice("y el aviso viaja al Excel, no solo a la consola", _txt, "no se suman jamás")
    c.dice("NO borra al cliente MARIA TESTA", out, "MARIA TESTA", debe=False)
    c.val("transportadoras listadas (sin umbral)", len(filas(L, "POR TRANSPORTADORA")), 4)
    c.val("ciudades listadas (sin umbral)", len(filas(L, "MEJOR TRANSP x CIUDAD")), 4)
    c.val("departamentos listados (sin umbral)", len(filas(L, "POR DEPARTAMENTO")), 4)
    c.dice("hoja POSIBLES FANTASMA", ",".join(openpyxl.load_workbook(L).sheetnames), "POSIBLES FANTASMA")
    if hay_pdf:
        c.dice("genera el PDF", out, "RESUMEN_EJECUTIVO.pdf")
        _pdf = os.path.join(raiz, "banco", "Analisis", "RESUMEN_EJECUTIVO.pdf")
        c.val("el PDF existe", os.path.exists(_pdf), True)
        # Comprobar solo que el archivo existe no prueba NADA de lo que dice por dentro.
        _b = io.open(_pdf, "rb").read()
        c.val("el PDF no sale vacio", len(_b) > 1500, True)
        _t = _texto_pdf(_pdf)
        if _t is None:
            c.omitir("contenido del PDF", "no hay con que extraer el texto en este Python")
        else:
            c.dice("el PDF lleva la cifra correcta de ganancia", _t, "225.000")
            c.dice("el PDF declara lo que quedo fuera", _t, "fuera de estas cifras")
            c.dice("el PDF avisa de que suma dos cuentas", _t.lower(), "cuentas")
            c.dice("los acentos salen bien en el render", _t, "Dónde")
    else:
        c.omitir("PDF", "sin reportlab en este Python")
        c.dice("avisa de que falta reportlab", out, "sin 'reportlab'")
        c.dice("aun asi entrega LOGISTICA", out, "MAESTRO_LOGISTICA.xlsx")
        c.dice("aun asi entrega CONTACTOS", out, "MAESTRO_CONTACTOS.xlsx")

    print("\n=== C1 control POSITIVO: pruebas de verdad y export duplicado ===")
    rc, out = correr(os.path.join(raiz, "C1"))
    c.val("codigo de salida", rc, 0)
    c.dice("sigue cazando PRUEBA JUAN", out, "PRUEBA JUAN")
    c.dice("sigue cazando TEST", out, "TEST (3009999998)")
    c.dice("no dobla el conteo de excluidos", out, "EXCLUIDOS por ser prueba: 2")
    c.dice("caza el export duplicado", out, "Duplicados detectados: 2")
    c.val("ganancia sin las pruebas",
          celda(os.path.join(raiz, "C1", "Analisis", "MAESTRO_LOGISTICA.xlsx"),
                "RESUMEN EJECUTIVO", "Ganancia realizada (de lo entregado)"), "$20.000")

    print("\n=== C2: falta el informe POR PEDIDO ===")
    rc, out = correr(os.path.join(raiz, "C2"))
    c.val("codigo de salida", rc, 0)
    v = celda(os.path.join(raiz, "C2", "Analisis", "MAESTRO_LOGISTICA.xlsx"), "RESUMEN EJECUTIVO", "VEREDICTO")
    c.val("veredicto bloqueado", v, "SIN VEREDICTO: FALTA EL INFORME POR PEDIDO")
    c.dice("no inventa un NO RENTABLE", str(v), "NO RENTABLE", debe=False)

    print("\n=== C3: config completa ===")
    rc, out = correr(os.path.join(raiz, "C3"))
    c.val("codigo de salida", rc, 0)
    L3 = os.path.join(raiz, "C3", "Analisis", "MAESTRO_LOGISTICA.xlsx")
    c.val("titulo con la marca del CLIENTE",
          openpyxl.load_workbook(L3, data_only=True)["RESUMEN EJECUTIVO"]["B2"].value,
          "RESUMEN EJECUTIVO — DISTRIBUIDORA X · DROPI")
    c.val("veredicto", celda(L3, "RESUMEN EJECUTIVO", "VEREDICTO"), "RENTABLE")
    c.val("utilidad neta 100k-10k-30k", celda(L3, "RESUMEN EJECUTIVO", "(=) UTILIDAD NETA FINAL"), "$60.000")

    print("\n=== C5: ninguna orden cerrada todavia ===")
    rc, out = correr(os.path.join(raiz, "C5"))
    c.val("codigo de salida", rc, 0)
    p5 = str(celda(os.path.join(raiz, "C5", "Analisis", "MAESTRO_LOGISTICA.xlsx"),
                   "RESUMEN", "% ENTREGA"))
    c.dice("no dice 0,0% sin nada cerrado", p5, "0.0%", debe=False)
    c.dice("dice que aun no se sabe", p5, "s/d")
    v5 = str(celda(os.path.join(raiz, "C5", "Analisis", "MAESTRO_LOGISTICA.xlsx"),
                   "RESUMEN EJECUTIVO", "VEREDICTO"))
    c.dice("NO dictamina NO RENTABLE sin una sola orden cerrada", v5, "NO RENTABLE", debe=False)
    c.dice("dice que no hay veredicto todavía", v5, "SIN VEREDICTO")

    print("\n=== C6: ENTREGADO A TRANSPORTADORA no es una venta cobrada ===")
    rc, out = correr(os.path.join(raiz, "C6"))
    c.val("codigo de salida", rc, 0)
    L6 = os.path.join(raiz, "C6", "Analisis", "MAESTRO_LOGISTICA.xlsx")
    c.val("solo 1 entregada de 3", celda(L6, "RESUMEN", "Entregadas"), 1)
    c.val("las otras 2 van en transito", celda(L6, "RESUMEN", "En tránsito"), 2)
    c.val("la ganancia cobrada es solo la entregada",
          celda(L6, "RESUMEN EJECUTIVO", "Ganancia realizada (de lo entregado)"), "$20.000")
    c.val("y el RESUMEN dice lo mismo que el RESUMEN EJECUTIVO",
          celda(L6, "RESUMEN", "Ganancia cobrada (solo entregadas)"), "$20.000")

    print("\n=== C7: basura que no puede convertirse en cifras falsas ===")
    rc, out = correr(os.path.join(raiz, "C7"))
    c.val("codigo de salida", rc, 0)
    L7 = os.path.join(raiz, "C7", "Analisis", "MAESTRO_LOGISTICA.xlsx")
    c.val("las 2 filas sin ID NO se funden en una", celda(L7, "RESUMEN", "Órdenes totales"), 5)
    c.dice("declara las filas sin ID", out, "Filas SIN columna ID: 2")
    # 20.000 + 20.000 + 80.000 (8E+04) + (-20.000) + 10.000 = 110.000
    c.val("el dinero raro se lee bien",
          celda(L7, "RESUMEN EJECUTIVO", "Ganancia realizada (de lo entregado)"), "$110.000")
    c.dice("declara la fecha futura", out, "Fechas sospechosas")
    c.dice("declara la CANTIDAD rara", out, "CANTIDAD en cero o negativa")
    _prods = filas(L7, "POR PRODUCTO")
    c.dice("el producto sin nombre sale como (SIN DATO)", ",".join(map(str, _prods)), "(SIN DATO)")

    print("\n=== C8: falta el export POR PRODUCTO ===")
    rc, out = correr(os.path.join(raiz, "C8"))
    c.val("codigo de salida", rc, 0)
    c.dice("avisa de que la base de clientes no se pudo medir", out, "POR PRODUCTO")
    c.dice("y explica que el 0 no es 'no tienes clientes'", out, "no se pudo medir")
    _hojas8 = openpyxl.load_workbook(os.path.join(raiz, "C8", "Analisis", "MAESTRO_LOGISTICA.xlsx")).sheetnames
    c.dice("el Excel trae hoja COBERTURA", ",".join(_hojas8), "COBERTURA")

    print("\n=== C4 control NEGATIVO: compras reales, no fantasmas ===")
    rc, out = correr(os.path.join(raiz, "C4"))
    c.val("codigo de salida", rc, 0)
    c.dice("no inventa fantasmas", out, "fantasma: 0 par(es)")

# --------------------------------------------------------------------------- mutacion
MUTACIONES = [
 ("classify vuelve a ignorar SINIESTRO",
  'or "SINIESTRO" in e or "PERDID" in e', 'or "PERDID" in e'),
 ("money vuelve a multiplicar por 100",
  '        s = (p[0] + "." + p[1]) if (len(p) == 2 and 1 <= len(p[1]) <= 2) else "".join(p)',
  '        s = "".join(p)'),
 ("es_prueba vuelve a excluir por subcadena",
  'r"(?<![0-9A-ZÁÉÍÓÚÑ])(?:" + alt +\n'
  '                                 r")(?![0-9A-ZÁÉÍÓÚÑ])"',
  'r"(?:" + alt + r")"'),
 ("classify vuelve a ignorar GUIA_ANULADA",
  'if "CANCELAD" in e or "RECHAZAD" in e or "ANULAD" in e: return "cancelado"',
  'if e in ("CANCELADO", "RECHAZADO"): return "cancelado"'),
 ("vuelve el 0,0% enganoso sin nada cerrado",
  'return f"{x*100:.1f}%" if cerradas else "s/d (nada cerrado aún)"',
  'return f"{x*100:.1f}%"'),
 ("vuelve el NO RENTABLE falso sin ordenes cerradas",
  'elif _cerradas == 0:', 'elif False:'),
 # --- estas atacan REGLAS DEL ESTANDAR que antes no cubria ninguna mutacion ---
 ("ENTREGADO A TRANSPORTADORA pasa a contar como venta cobrada",
  'if "ENTREGADO A TRANSPORTADORA" in e: return "transito"',
  'if "ENTREGADO A TRANSPORTADORA" in e: return "entregado"'),
 ("las candidatas a fantasma se descuentan automaticamente",
  'if CAND:\n    print(', 'if CAND:\n    ped = [x for x in ped if x["oid"] not in {a["oid"] for a, b in CAND}]\n    print('),
 ("el ID vacio vuelve a fundir filas distintas",
  '_k = (_x["cuenta"], _x["oid"]) if _x["oid"] else ("__sin_id__", _i)',
  '_k = (_x["cuenta"], _x["oid"])'),
 ("la ganancia del RESUMEN vuelve a sumar lo no cobrado",
  'gan = sum(x["ganancia"] for x in recs if x["clase"] == "entregado")',
  'gan = sum(x["ganancia"] for x in recs)'),
]

def mutar(raiz):
    """Rompe el motor a proposito y comprueba que la autoprueba lo caza.

    Trabaja sobre una COPIA en un temporal: antes mutaba el motor de produccion in situ
    dentro de ~/.claude/skills, que es donde el sandbox deniega escritura (reventaba con
    PermissionError), y un corte a mitad del bucle dejaba la skill instalada mutada."""
    global MOTOR
    print("\n" + "=" * 70)
    print("PRUEBA DEL INSTRUMENTO: se rompe el motor a proposito, la autoprueba debe cazarlo")
    print("=" * 70)
    original = MOTOR
    tmp = tempfile.mkdtemp(prefix="autoprueba-mutar-")
    copia = os.path.join(tmp, "motor_analisis_dropi.py")
    shutil.copy(original, copia)
    fuente = io.open(original, encoding="utf-8").read()
    sin_cazar = []
    try:
        MOTOR = copia
        for nombre, viejo, nuevo in MUTACIONES:
            if fuente.count(viejo) != 1:
                print(f"  ??      {nombre}: no se pudo inyectar ({fuente.count(viejo)} coincidencias)")
                sin_cazar.append(nombre + " [no inyectada]"); continue
            io.open(copia, "w", encoding="utf-8").write(fuente.replace(viejo, nuevo))
            cc = Chequeo()
            import contextlib
            with contextlib.redirect_stdout(io.StringIO()):
                probar(raiz, cc)
            if cc.fallos: print(f"  CAZADA  {nombre}  ({len(cc.fallos)} comprobacion/es en rojo)")
            else: print(f"  NO CAZADA <<< {nombre}"); sin_cazar.append(nombre)
    finally:
        MOTOR = original
        shutil.rmtree(tmp, ignore_errors=True)
        print("el motor instalado NO se tocó en ningún momento")
    return sin_cazar

def main():
    if not os.path.exists(MOTOR): sys.exit(f"No encuentro el motor en {MOTOR}")
    raiz = tempfile.mkdtemp(prefix="autoprueba-dropi-")
    try:
        sembrar(raiz)
        c = Chequeo()
        probar(raiz, c)
        sin_cazar = mutar(raiz) if "--mutar" in sys.argv else None
        print("\n" + "=" * 70)
        print(f"COBERTURA: {c.ok} comprobacion(es) en verde, {len(c.fallos)} en rojo"
              + (f", {len(c.omitidos)} omitida(s)" if c.omitidos else ""))
        for f in c.fallos: print("  FALLA   ", f)
        for o in c.omitidos: print("  OMITIDO ", o)
        if sin_cazar is not None:
            print(f"MUTACIONES no cazadas: {len(sin_cazar)}")
            for m in sin_cazar: print("  ", m)
        malo = bool(c.fallos) or bool(sin_cazar)
        print("RESULTADO:", "HAY FALLAS" if malo else "todas las comprobaciones cuadran")
        return 1 if malo else 0
    finally:
        shutil.rmtree(raiz, ignore_errors=True)

if __name__ == "__main__":
    sys.exit(main())
